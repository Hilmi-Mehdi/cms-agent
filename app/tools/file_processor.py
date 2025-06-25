"""File processor tool for extracting content from various file formats."""

import os
import asyncio
from typing import Dict, Any, List
from pathlib import Path

import aiofiles
from pptx import Presentation
from docx import Document
import pypdf

from app.tools.base_tool import BaseTool, ToolDefinition, ToolParameter
from app.models.schemas import ToolResult
from app.config import settings


class FileProcessorTool(BaseTool):
    """Tool for processing and extracting content from course files."""
    
    def get_definition(self) -> ToolDefinition:
        """Return the tool definition."""
        return ToolDefinition(
            name="file_processor",
            description="Extract and process content from course files (PDF, PPTX, DOCX, TXT, MD)",
            parameters=[
                ToolParameter(
                    name="file_path",
                    type="string",
                    description="Path to the file to process",
                    required=True
                ),
                ToolParameter(
                    name="extraction_type",
                    type="string",
                    description="Type of content extraction",
                    required=False,
                    default="text",
                    enum=["text", "structured", "metadata", "full"]
                ),
                ToolParameter(
                    name="include_images",
                    type="boolean",
                    description="Whether to include image descriptions (where possible)",
                    required=False,
                    default=False
                ),
                ToolParameter(
                    name="max_size_mb",
                    type="number",
                    description="Maximum file size in MB to process",
                    required=False,
                    default=50
                )
            ]
        )
    
    async def execute(self, parameters: Dict[str, Any]) -> ToolResult:
        """Execute file processing."""
        file_path = parameters["file_path"]
        extraction_type = parameters.get("extraction_type", "text")
        include_images = parameters.get("include_images", False)
        max_size_mb = parameters.get("max_size_mb", 50)
        
        try:
            # Validate file
            validation_result = await self._validate_file(file_path, max_size_mb)
            if not validation_result["valid"]:
                return ToolResult(
                    success=False,
                    error=validation_result["error"]
                )
            
            file_info = validation_result["file_info"]
            
            # Extract content based on file type
            content_result = await self._extract_content(
                file_path, file_info["extension"], extraction_type, include_images
            )
            
            return ToolResult(
                success=True,
                data={
                    "file_info": file_info,
                    "content": content_result,
                    "extraction_type": extraction_type
                },
                suggested_next_tools=["content_analyzer"],
                agent_notes=f"Processed {file_info['size_mb']:.2f}MB {file_info['extension']} file"
            )
            
        except Exception as e:
            return ToolResult(
                success=False,
                error=f"File processing failed: {str(e)}"
            )
    
    async def _validate_file(self, file_path: str, max_size_mb: float) -> Dict[str, Any]:
        """Validate file exists, size, and type."""
        try:
            file_path_obj = Path(file_path)
            
            if not file_path_obj.exists():
                return {"valid": False, "error": f"File not found: {file_path}"}
            
            if not file_path_obj.is_file():
                return {"valid": False, "error": f"Path is not a file: {file_path}"}
            
            file_size = file_path_obj.stat().st_size
            size_mb = file_size / (1024 * 1024)
            
            if size_mb > max_size_mb:
                return {"valid": False, "error": f"File too large: {size_mb:.2f}MB > {max_size_mb}MB"}
            
            extension = file_path_obj.suffix.lower()
            
            if extension not in settings.supported_file_types:
                return {
                    "valid": False, 
                    "error": f"Unsupported file type: {extension}. Supported: {settings.supported_file_types}"
                }
            
            return {
                "valid": True,
                "file_info": {
                    "name": file_path_obj.name,
                    "extension": extension,
                    "size_bytes": file_size,
                    "size_mb": size_mb,
                    "path": str(file_path_obj.absolute())
                }
            }
            
        except Exception as e:
            return {"valid": False, "error": f"File validation error: {str(e)}"}
    
    async def _extract_content(
        self, 
        file_path: str, 
        extension: str, 
        extraction_type: str, 
        include_images: bool
    ) -> Dict[str, Any]:
        """Extract content based on file type."""
        extractors = {
            ".txt": self._extract_text_file,
            ".md": self._extract_text_file,
            ".pdf": self._extract_pdf,
            ".pptx": self._extract_pptx,
            ".docx": self._extract_docx
        }
        
        extractor = extractors.get(extension)
        if not extractor:
            raise ValueError(f"No extractor available for {extension}")
        
        return await extractor(file_path, extraction_type, include_images)
    
    async def _extract_text_file(
        self, 
        file_path: str, 
        extraction_type: str, 
        include_images: bool
    ) -> Dict[str, Any]:
        """Extract content from text files."""
        async with aiofiles.open(file_path, 'r', encoding='utf-8') as file:
            content = await file.read()
        
        result = {
            "text": content,
            "metadata": {
                "line_count": len(content.splitlines()),
                "char_count": len(content),
                "word_count": len(content.split())
            }
        }
        
        if extraction_type in ["structured", "full"]:
            # Try to detect structure in markdown files
            if file_path.endswith('.md'):
                result["structure"] = self._parse_markdown_structure(content)
        
        return result
    
    async def _extract_pdf(
        self, 
        file_path: str, 
        extraction_type: str, 
        include_images: bool
    ) -> Dict[str, Any]:
        """Extract content from PDF files."""
        def extract_pdf_content():
            with open(file_path, 'rb') as file:
                reader = pypdf.PdfReader(file)
                
                text_content = []
                metadata = {
                    "page_count": len(reader.pages),
                    "title": reader.metadata.get('/Title', '') if reader.metadata else '',
                    "author": reader.metadata.get('/Author', '') if reader.metadata else ''
                }
                
                for page_num, page in enumerate(reader.pages):
                    page_text = page.extract_text()
                    text_content.append({
                        "page": page_num + 1,
                        "text": page_text
                    })
                
                return {
                    "text": "\n\n".join([page["text"] for page in text_content]),
                    "pages": text_content if extraction_type in ["structured", "full"] else None,
                    "metadata": metadata
                }
        
        # Run PDF extraction in thread pool
        return await asyncio.get_event_loop().run_in_executor(None, extract_pdf_content)
    
    async def _extract_pptx(
        self, 
        file_path: str, 
        extraction_type: str, 
        include_images: bool
    ) -> Dict[str, Any]:
        """Extract content from PowerPoint files."""
        def extract_pptx_content():
            presentation = Presentation(file_path)
            
            slides_content = []
            all_text = []
            
            for slide_num, slide in enumerate(presentation.slides):
                slide_text = []
                
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_text.append(shape.text)
                
                slide_content = {
                    "slide": slide_num + 1,
                    "text": "\n".join(slide_text)
                }
                
                slides_content.append(slide_content)
                all_text.append(slide_content["text"])
            
            return {
                "text": "\n\n".join(all_text),
                "slides": slides_content if extraction_type in ["structured", "full"] else None,
                "metadata": {
                    "slide_count": len(presentation.slides),
                    "title": getattr(presentation.core_properties, 'title', '') or ''
                }
            }
        
        return await asyncio.get_event_loop().run_in_executor(None, extract_pptx_content)
    
    async def _extract_docx(
        self, 
        file_path: str, 
        extraction_type: str, 
        include_images: bool
    ) -> Dict[str, Any]:
        """Extract content from Word documents."""
        def extract_docx_content():
            doc = Document(file_path)
            
            paragraphs = []
            all_text = []
            
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraph_data = {
                        "text": para.text,
                        "style": para.style.name if para.style else "Normal"
                    }
                    paragraphs.append(paragraph_data)
                    all_text.append(para.text)
            
            return {
                "text": "\n\n".join(all_text),
                "paragraphs": paragraphs if extraction_type in ["structured", "full"] else None,
                "metadata": {
                    "paragraph_count": len(paragraphs),
                    "title": doc.core_properties.title or '',
                    "author": doc.core_properties.author or ''
                }
            }
        
        return await asyncio.get_event_loop().run_in_executor(None, extract_docx_content)
    
    def _parse_markdown_structure(self, content: str) -> Dict[str, Any]:
        """Parse markdown structure."""
        lines = content.split('\n')
        structure = {
            "headers": [],
            "sections": []
        }
        
        current_section = None
        
        for line_num, line in enumerate(lines):
            stripped = line.strip()
            
            # Detect headers
            if stripped.startswith('#'):
                level = len(stripped) - len(stripped.lstrip('#'))
                header_text = stripped.lstrip('#').strip()
                
                header = {
                    "level": level,
                    "text": header_text,
                    "line": line_num + 1
                }
                
                structure["headers"].append(header)
                
                # Start new section
                if current_section:
                    structure["sections"].append(current_section)
                
                current_section = {
                    "title": header_text,
                    "level": level,
                    "start_line": line_num + 1,
                    "content": []
                }
            
            elif current_section and stripped:
                current_section["content"].append(stripped)
        
        # Add last section
        if current_section:
            structure["sections"].append(current_section)
        
        return structure 